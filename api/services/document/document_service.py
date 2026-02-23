"""
Document Service - Handles document generation (DOCX, PDF, Markdown).

This is the main facade that orchestrates document generation using
helper modules for format-specific generation and template rendering.
"""

from typing import Dict, List, Any, Tuple
from pathlib import Path
import os
import re
from datetime import datetime

from api.config import get_settings
from .document_generators import (
    generate_docx,
    generate_pdf,
    generate_markdown,
    generate_html,
    parse_markdown_table,
)
from .document_template_helpers import (
    render_template,
    flatten_dict,
    render_sections,
    categorize_skills,
    build_project_company_data,
    calculate_total_experience,
    extract_all_skills,
)

settings = get_settings()


class DocumentService:
    """Service for document generation and template parsing."""

    def __init__(self):
        self.result_dir = settings.result_dir

    async def parse_template(self, file_path: str) -> Tuple[str, Dict[str, str]]:
        """Parse a template file and extract field mappings."""
        ext = Path(file_path).suffix.lower()

        if ext in [".docx", ".doc"]:
            return await self._parse_docx_template(file_path)
        elif ext == ".pdf":
            return await self._parse_pdf_template(file_path)
        else:
            raise ValueError(f"Unsupported template format: {ext}")

    async def _parse_docx_template(self, file_path: str) -> Tuple[str, Dict[str, str]]:
        """Parse a Word document template."""
        from docx import Document

        doc = Document(file_path)
        content_parts = []
        field_mappings = {}

        # Pattern for placeholders like {{field_name}} or {field_name}
        placeholder_pattern = r"\{\{?\s*([a-zA-Z_][a-zA-Z0-9_\.]*)\s*\}?\}"

        for para in doc.paragraphs:
            text = para.text
            content_parts.append(text)

            # Find all placeholders
            matches = re.findall(placeholder_pattern, text)
            for match in matches:
                field_mappings[f"{{{{{match}}}}}"] = match

        full_content = "\n".join(content_parts)
        return full_content, field_mappings

    async def _parse_pdf_template(self, file_path: str) -> Tuple[str, Dict[str, str]]:
        """Parse a PDF template (extract text and find placeholders)."""
        from PyPDF2 import PdfReader

        reader = PdfReader(file_path)
        content_parts = []
        field_mappings = {}

        placeholder_pattern = r"\{\{?\s*([a-zA-Z_][a-zA-Z0-9_\.]*)\s*\}?\}"

        for page in reader.pages:
            text = page.extract_text()
            content_parts.append(text)

            matches = re.findall(placeholder_pattern, text)
            for match in matches:
                field_mappings[f"{{{{{match}}}}}"] = match

        full_content = "\n".join(content_parts)
        return full_content, field_mappings

    async def generate_document(
        self,
        template_content: str,
        data: Dict[str, Any],
        output_format: str,
        document_name: str,
        template_platform: str = None,
    ) -> Tuple[str, int]:
        """Generate a document from template and data.

        Args:
            template_content: Template content (markdown/mustache format)
            data: Data to fill into template
            output_format: Output format (docx, pdf, md, html)
            document_name: Name for the output file
            template_platform: Platform type for special handling
        """
        # Generate unique filename
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_name = re.sub(r"[^\w\-_]", "_", document_name)
        filename = f"{safe_name}_{timestamp}.{output_format}"
        file_path = self.result_dir / filename

        os.makedirs(self.result_dir, exist_ok=True)

        # Use specialized generator for Korean resume templates
        if output_format == "docx" and template_platform in (
            "career_description",
            "career_description_no_personal",
            "resume",
        ):
            from api.services.docx.docx_generator import DocxGenerator

            generator = DocxGenerator()

            if template_platform == "career_description":
                file_size = generator.generate_career_description(
                    data, file_path, include_personal_info=True
                )
            elif template_platform == "career_description_no_personal":
                file_size = generator.generate_career_description(
                    data, file_path, include_personal_info=False
                )
            elif template_platform == "resume":
                file_size = generator.generate_resume(data, file_path)
            else:
                file_size = generator.generate_career_description(data, file_path)

            return str(file_path), file_size

        # For other templates, use markdown-based rendering
        rendered_content = self._render_template(template_content, data)

        if output_format == "docx":
            file_size = await generate_docx(rendered_content, file_path)
        elif output_format == "pdf":
            file_size = await generate_pdf(rendered_content, file_path)
        elif output_format == "md":
            file_size = await generate_markdown(rendered_content, file_path)
        elif output_format == "html":
            file_size = await generate_html(rendered_content, file_path)
        else:
            raise ValueError(f"Unsupported output format: {output_format}")

        return str(file_path), file_size

    def _render_template(self, template: str, data: Dict[str, Any]) -> str:
        """Render template with data (simple mustache-like syntax)."""
        return render_template(template, data)

    def _flatten_dict(self, d: Dict, parent_key: str = "", sep: str = ".") -> Dict:
        """Flatten nested dictionary."""
        return flatten_dict(d, parent_key, sep)

    def _render_sections(self, template: str, data: Dict[str, Any]) -> str:
        """Render section loops in template."""
        return render_sections(template, data)

    def _parse_markdown_table(
        self, lines: List[str], start_idx: int
    ) -> Tuple[List[List[str]], int]:
        """Parse markdown table starting at given index."""
        return parse_markdown_table(lines, start_idx)

    def prepare_template_data(
        self,
        user_data: Dict[str, Any],
        companies: List[Dict[str, Any]],
        projects: List[Dict[str, Any]],
    ) -> Dict[str, Any]:
        """Prepare data structure for template rendering."""
        # Extract categorized skills from all projects
        skill_categories = categorize_skills(projects)

        return {
            # Basic user info
            "name": user_data.get("name", ""),
            "email": user_data.get("email", ""),
            "github_username": user_data.get("github_username", ""),
            "summary": user_data.get("summary", ""),
            # Personal information (from user profile)
            "birthdate": user_data.get("birthdate", ""),
            "address": user_data.get("address", ""),
            "phone": user_data.get("phone", ""),
            "military_status": user_data.get("military_status", ""),
            "military_type": user_data.get("military_type", ""),
            # Education
            "university": user_data.get("university", ""),
            "major": user_data.get("major", ""),
            "graduation_date": user_data.get("graduation_date", ""),
            "gpa": user_data.get("gpa", ""),
            # Categorized skills
            "programming_languages": skill_categories.get("programming_languages", ""),
            "frameworks": skill_categories.get("frameworks", ""),
            "databases": skill_categories.get("databases", ""),
            "tools": skill_categories.get("tools", ""),
            "cloud": skill_categories.get("cloud", ""),
            "devops": skill_categories.get(
                "cloud", ""
            ),  # Alias for templates using devops
            "tooling": skill_categories.get(
                "tools", ""
            ),  # Alias for templates using tooling
            # Legacy skills (all combined)
            "skills": extract_all_skills(projects),
            # Salary and job change info (from user profile)
            "current_salary": user_data.get("current_salary", ""),
            "desired_salary": user_data.get("desired_salary", ""),
            "job_change_reason": user_data.get("job_change_reason", ""),
            "available_date": user_data.get("available_date", ""),
            "motivation": user_data.get("motivation", ""),
            "competencies": user_data.get("competencies", ""),
            "personality": user_data.get("personality", ""),
            # Companies/Career
            "companies": [
                {
                    "name": c.get("name", ""),
                    "position": c.get("position", ""),
                    "department": c.get("department", ""),
                    "start_date": str(c.get("start_date", "")),
                    "end_date": str(c.get("end_date", ""))
                    if c.get("end_date")
                    else "현재",
                    "description": c.get("description", ""),
                    "location": c.get("location", ""),
                }
                for c in companies
            ],
            # Build company lookup for project enrichment
            **build_project_company_data(projects, companies),
            # Total career experience
            "total_experience": calculate_total_experience(companies),
        }

    def _categorize_skills(self, projects: List[Dict[str, Any]]) -> Dict[str, str]:
        """Categorize skills from projects into predefined categories."""
        return categorize_skills(projects)

    def _build_project_company_data(
        self, projects: List[Dict[str, Any]], companies: List[Dict[str, Any]]
    ) -> Dict[str, Any]:
        """Build project data enriched with company info."""
        return build_project_company_data(projects, companies)

    def _calculate_total_experience(self, companies: List[Dict[str, Any]]) -> str:
        """Calculate total career experience from companies."""
        return calculate_total_experience(companies)

    def _extract_all_skills(self, projects: List[Dict[str, Any]]) -> str:
        """Extract and deduplicate all skills from projects."""
        return extract_all_skills(projects)
